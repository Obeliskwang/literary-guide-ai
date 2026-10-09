import os
import tempfile
from typing import Optional

def parse_document(file_bytes: bytes, filename: str) -> Optional[str]:
    """统一的文档解析入口，根据扩展名自动选择解析器"""
    ext = filename.lower().split(".")[-1]
    
    with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as f:
        f.write(file_bytes)
        temp_path = f.name
    
    try:
        if ext == "txt":
            return parse_txt(temp_path)
        elif ext == "docx":
            return parse_docx(temp_path)
        elif ext == "pdf":
            return parse_pdf(temp_path)
        else:
            return None
    finally:
        os.unlink(temp_path)

def parse_txt(file_path: str) -> str:
    """读取 TXT 文件，自动检测编码"""
    encodings = ["utf-8", "gbk", "gb2312", "utf-16"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码，请使用 UTF-8 或 GBK 格式")

def parse_docx(file_path: str) -> str:
    """读取 Word 文档"""
    try:
        from docx import Document
    except ImportError:
        raise ImportError("请安装 python-docx：pip install python-docx")
    
    doc = Document(file_path)
    text_parts = []
    
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)
    
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                text_parts.append(" | ".join(row_text))
    
    return "\n".join(text_parts)

def parse_pdf(file_path: str) -> str:
    """读取 PDF 文档，优先使用 pdfplumber，降级使用 PyPDF2"""
    try:
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            text = ""
            for page in pdf.pages:
                page_text = page.extract_text(layout=True)
                if page_text:
                    text += page_text + "\n\n"
            return text
    except ImportError:
        pass
    
    try:
        import PyPDF2
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n\n"
            return text
    except ImportError:
        raise ImportError("请安装 pdfplumber 或 PyPDF2：pip install pdfplumber PyPDF2")
